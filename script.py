# import os
# import re
# import json
# import argparse
# import logging
# import requests
# import io
# from bs4 import BeautifulSoup
# from docx import Document
# from docx.shared import Inches
# from collections import OrderedDict
# from PIL import Image

# # ---------------------------
# # Logging
# # ---------------------------
# logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")


# # ---------------------------
# # Helpers
# # ---------------------------
# def clean_whitespace(s: str) -> str:
#     if not s:
#         return s
#     return " ".join(s.split())


# def dedup_preserve_order(items):
#     seen = set()
#     out = []
#     for it in items:
#         it = it.strip()
#         if not it or it in seen:
#             continue
#         seen.add(it)
#         out.append(it)
#     return out


# def sanitize_filename(name):
#     return re.sub(r'[\\/*?:"<>|]', "", name)


# # ---------------------------
# # Image download logic
# # ---------------------------
# def download_image(image_url, folder_path, image_name):
#     """Download and save an image as JPG with resizing"""
#     if not image_url or not image_url.startswith("http"):
#         image_url = "https:" + image_url if image_url and image_url.startswith("//") else None
#         if not image_url:
#             return None
#     try:
#         response = requests.get(image_url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
#         response.raise_for_status()
#         img = Image.open(io.BytesIO(response.content))
#         # resize small images to min 650x650
#         if img.size[0] < 640 or img.size[1] < 640:
#             resample_filter = getattr(Image, "Resampling", Image).LANCZOS
#             img = img.resize((650, 650), resample_filter)
#         if img.mode != "RGB":
#             img = img.convert("RGB")
#         safe_image_name = sanitize_filename(image_name)
#         filename = f"{safe_image_name}.jpg"
#         os.makedirs(folder_path, exist_ok=True)
#         filepath = os.path.join(folder_path, filename)
#         img.save(filepath, "JPEG", quality=95)
#         return filepath
#     except Exception as e:
#         logging.warning(f"Failed to download image {image_url}. Error: {e}")
#         return None


# def save_product_images(product_name, product_id, image_urls):
#     """Download all images for a product into a dedicated folder"""
#     folder = os.path.join(os.getcwd(), "product_images", sanitize_filename(product_id or product_name))
#     os.makedirs(folder, exist_ok=True)
#     saved = []
#     for idx, url in enumerate(image_urls):
#         img_path = download_image(url, folder, f"{product_id}_{idx+1}")
#         if img_path:
#             saved.append(img_path)
#     return saved


# # ---------------------------
# # JSON-LD parsing
# # ---------------------------
# def parse_json_ld(soup):
#     results = {}
#     for script in soup.find_all("script", type="application/ld+json"):
#         try:
#             raw = script.string
#             if not raw:
#                 continue
#             data = json.loads(raw.strip())
#         except Exception:
#             try:
#                 raw_text = "".join(script.contents)
#                 data = json.loads(raw_text.strip())
#             except Exception:
#                 continue
#         if isinstance(data, list):
#             for entry in data:
#                 if isinstance(entry, dict) and entry.get("@type", "").lower() in ("product", "productmodel"):
#                     data = entry
#                     break
#         if isinstance(data, dict) and data.get("@type", "").lower() in ("product", "productmodel"):
#             results["name"] = data.get("name") or results.get("name")
#             results["description"] = data.get("description") or results.get("description")
#             offers = data.get("offers") or {}
#             if isinstance(offers, list):
#                 offers = offers[0]
#             if isinstance(offers, dict):
#                 price = offers.get("price") or offers.get("priceSpecification", {}).get("price")
#                 currency = offers.get("priceCurrency")
#                 if price:
#                     results["price"] = (str(price) if not currency else f"{price} {currency}")
#             return results
#     return results


# # ---------------------------
# # Extraction routines
# # ---------------------------
# PRICE_RE = re.compile(r"(?:₹|Rs\.?|INR)\s?[\d,]+(?:\.\d+)?", flags=re.IGNORECASE)


# def clean_price(s):
#     if s is None:
#         return None
#     s = str(s)
#     s = s.replace("Rs.", "").replace("Rs", "").replace("INR", "").replace("₹", "").strip()
#     m = re.search(r"[\d,]+(?:\.\d+)?", s)
#     if m:
#         return m.group(0)
#     return s


# def find_price(soup, page_text, jsonld):
#     if jsonld.get("price"):
#         return clean_price(jsonld["price"])
#     for meta_name in ("product:price:amount", "og:price:amount", "price"):
#         m = soup.find("meta", {"property": meta_name}) or soup.find("meta", {"name": meta_name})
#         if m and m.get("content"):
#             return clean_price(m.get("content"))
#     item_price = soup.find(attrs={"itemprop": "price"})
#     if item_price:
#         return clean_price(item_price.get("content") or item_price.get_text())
#     candidate = soup.find(lambda tag: tag.name in ("span", "div") and tag.get("class") and any("price" in c.lower() for c in tag.get("class")))
#     if candidate:
#         return clean_price(candidate.get_text(" ", strip=True))
#     match = PRICE_RE.search(page_text)
#     if match:
#         return clean_price(match.group(0))
#     return None


# def find_title(soup, jsonld):
#     if jsonld.get("name"):
#         return clean_whitespace(jsonld["name"])
#     h1 = soup.find("h1")
#     if h1 and h1.get_text(strip=True):
#         return clean_whitespace(h1.get_text(" ", strip=True))
#     if soup.title and soup.title.string:
#         return clean_whitespace(soup.title.string)
#     return None


# def find_description(soup, page_text, jsonld):
#     if jsonld.get("description"):
#         return clean_whitespace(jsonld["description"])
#     desc_elem = soup.find(attrs={"itemprop": "description"})
#     if desc_elem and desc_elem.get_text(strip=True):
#         return clean_whitespace(desc_elem.get_text(" ", strip=True))
#     meta = soup.find("meta", {"name": "description"})
#     if meta and meta.get("content"):
#         return clean_whitespace(meta.get("content"))
#     paragraphs = [p.get_text(" ", strip=True) for p in soup.find_all("p")]
#     paragraphs = [clean_whitespace(p) for p in paragraphs if p and len(p) > 30]
#     if paragraphs:
#         paragraphs.sort(key=lambda s: len(s), reverse=True)
#         return paragraphs[0]
#     return None


# def extract_features(soup, page_text, jsonld):
#     features = []
#     if jsonld.get("additional_properties"):
#         features.extend(jsonld["additional_properties"])
#     # colon separated lines fallback
#     for line in page_text.splitlines():
#         line = line.strip()
#         if ":" in line and len(line) < 120:
#             left, right = line.split(":", 1)
#             if len(left) < 40 and len(right) > 1:
#                 features.append(f"{left.strip()}: {right.strip()}")
#     features = dedup_preserve_order(features)
#     return features


# def extract_images(soup, jsonld=None):
#     images = []
#     for img in soup.find_all("img"):
#         src = img.get("src") or img.get("data-src")
#         if src:
#             if src.startswith("//"):
#                 src = "https:" + src
#             images.append(src.split("?")[0])
#     return dedup_preserve_order(images)


# # ---------------------------
# # Main processor
# # ---------------------------
# def process_html_folder_to_docx(folder_path, output_docx="products.docx"):
#     html_files = []
#     for root, _, files in os.walk(folder_path):
#         for fname in files:
#             if fname.lower().endswith((".html", ".htm")):
#                 html_files.append(os.path.join(root, fname))

#     if not html_files:
#         logging.warning("No HTML files found in %s", folder_path)
#         return

#     doc = Document()
#     doc.add_heading("Extracted Product Details", level=0)

#     for file_path in sorted(html_files):
#         with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
#             html = f.read()

#         soup = BeautifulSoup(html, "lxml")
#         page_text = soup.get_text("\n", strip=True)

#         jsonld = parse_json_ld(soup)

#         title = find_title(soup, jsonld) or os.path.basename(file_path)
#         price = find_price(soup, page_text, jsonld) or "Not found"
#         description = find_description(soup, page_text, jsonld) or "Not found"
#         features = extract_features(soup, page_text, jsonld) or []

#         images = extract_images(soup, jsonld=jsonld)
#         saved_images = save_product_images(title, os.path.splitext(os.path.basename(file_path))[0], images)

#         doc.add_heading(title, level=1)
#         doc.add_paragraph(f"Source file: {os.path.basename(file_path)}")
#         doc.add_paragraph(f"Price: {price}")
#         doc.add_paragraph("Description:")
#         doc.add_paragraph(description)
#         if features:
#             doc.add_paragraph("Key Features:")
#             for feat in features:
#                 doc.add_paragraph(feat, style="List Bullet")

#         if saved_images:
#             doc.add_paragraph("Images:")
#             try:
#                 doc.add_picture(saved_images[0], width=Inches(2.5))
#             except Exception as e:
#                 logging.warning(f"Could not add image to docx: {e}")

#         doc.add_paragraph("")

#     doc.save(output_docx)
#     logging.info("Saved combined docx: %s", output_docx)


# # ---------------------------
# # CLI
# # ---------------------------
# def main():
#     parser = argparse.ArgumentParser(description="Extract product details + images from HTML files and write to DOCX.")
#     parser.add_argument("--input-dir", "-i", required=True, help="Folder containing HTML files")
#     parser.add_argument("--output", "-o", default="products.docx", help="Output .docx filename")
#     args = parser.parse_args()

#     if not os.path.isdir(args.input_dir):
#         logging.error("Input directory not found: %s", args.input_dir)
#         return

#     process_html_folder_to_docx(args.input_dir, output_docx=args.output)


# if __name__ == "__main__":
#     main()



# import os
# import re
# import json
# import argparse
# import logging
# import requests
# import io
# from bs4 import BeautifulSoup
# from docx import Document
# from docx.shared import Inches
# from collections import OrderedDict
# from PIL import Image
# from urllib.parse import urljoin, urlparse
# from datetime import datetime

# # ---------------------------
# # Logging
# # ---------------------------
# logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

# # ---------------------------
# # Helpers
# # ---------------------------
# def clean_whitespace(s: str) -> str:
#     """Removes extra whitespace from a string."""
#     if not s:
#         return s
#     return " ".join(s.split())

# def dedup_preserve_order(items):
#     """Deduplicates items in a list while preserving their original order."""
#     seen = set()
#     out = []
#     for it in items:
#         it = it.strip()
#         if not it or it in seen:
#             continue
#         seen.add(it)
#         out.append(it)
#     return out

# def sanitize_filename(name):
#     """Removes invalid characters from a string to make it a valid filename."""
#     return re.sub(r'[\\/*?:"<>|]', "", name)

# def get_base_url(soup, file_path):
#     """Extracts the base URL from the page to resolve relative URLs."""
#     base_tag = soup.find('base', href=True)
#     if base_tag and base_tag['href']:
#         return base_tag['href']
    
#     og_url = soup.find('meta', property='og:url')
#     if og_url and og_url.get('content'):
#         return og_url['content']
        
#     canonical = soup.find('link', rel='canonical')
#     if canonical and canonical.get('href'):
#         return canonical['href']
    
#     return f"file://{os.path.abspath(file_path)}"

# # ---------------------------
# # Image download logic
# # ---------------------------
# def download_image(image_url, folder_path, image_name):
#     """Download and save an image as JPG with resizing."""
#     if not image_url:
#         return None
#     try:
#         response = requests.get(image_url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
#         response.raise_for_status()
#         img = Image.open(io.BytesIO(response.content))
        
#         if img.size[0] < 640 or img.size[1] < 640:
#             resample_filter = getattr(Image, "Resampling", Image).LANCZOS
#             img = img.resize((650, 650), resample_filter)
        
#         if img.mode != "RGB":
#             img = img.convert("RGB")
            
#         safe_image_name = sanitize_filename(image_name)
#         filename = f"{safe_image_name}.jpg"
#         os.makedirs(folder_path, exist_ok=True)
#         filepath = os.path.join(folder_path, filename)
#         img.save(filepath, "JPEG", quality=95)
#         return filepath
#     except requests.exceptions.MissingSchema:
#         logging.warning(f"Skipping invalid image URL (missing http/https): {image_url}")
#         return None
#     except Exception as e:
#         logging.warning(f"Failed to download image {image_url}. Error: {e}")
#         return None

# def save_product_images(product_name, product_id, image_urls, base_output_path):
#     """Download all images for a product into a dedicated folder inside the base output path."""
#     if not product_name and not product_id:
#         logging.warning("Skipping image download due to missing product name/ID.")
#         return []
    
#     # Create product-specific folder inside the main timestamped output folder
#     product_folder = os.path.join(base_output_path, sanitize_filename(product_id or product_name))
#     os.makedirs(product_folder, exist_ok=True)
#     saved = []
    
#     for idx, url in enumerate(image_urls):
#         img_path = download_image(url, product_folder, f"{product_id}_{idx+1}")
#         if img_path:
#             saved.append(img_path)
#     return saved

# # ---------------------------
# # JSON-LD parsing
# # ---------------------------
# def parse_json_ld(soup):
#     """Parses JSON-LD scripts to find product information."""
#     results = {}
#     scripts = soup.find_all("script", type="application/ld+json")
#     for script in scripts:
#         try:
#             data = json.loads(script.string or "".join(script.contents))
#         except (json.JSONDecodeError, TypeError):
#             continue

#         if isinstance(data, list):
#             for item in data:
#                 if isinstance(item, dict) and item.get("@type", "").lower() in ("product", "productgroup", "productmodel"):
#                     data = item
#                     break
        
#         if isinstance(data, dict) and data.get("@type", "").lower() in ("product", "productgroup", "productmodel"):
#             results["name"] = data.get("name") or results.get("name")
#             results["description"] = data.get("description") or results.get("description")
            
#             image_data = data.get("image")
#             images = []
#             if isinstance(image_data, str):
#                 images.append(image_data)
#             elif isinstance(image_data, list):
#                 for img in image_data:
#                     if isinstance(img, str):
#                         images.append(img)
#                     elif isinstance(img, dict):
#                         images.append(img.get("url") or img.get("contentUrl"))
#             if images:
#                 results["images"] = images
            
#             offers = data.get("offers", {})
#             if isinstance(offers, list):
#                 offers = offers[0] if offers else {}
            
#             if isinstance(offers, dict):
#                 price = offers.get("price") or offers.get("priceSpecification", {}).get("price")
#                 currency = offers.get("priceCurrency")
#                 if price:
#                     results["price"] = (str(price) if not currency else f"{price} {currency}")
            
#             if "name" in results:
#                 return results
                
#     return results

# # ---------------------------
# # Extraction routines
# # ---------------------------
# PRICE_RE = re.compile(r"(?:₹|Rs\.?|INR)\s?[\d,]+(?:\.\d+)?", flags=re.IGNORECASE)

# def clean_price(s):
#     if s is None:
#         return None
#     s = str(s).replace("Rs.", "").replace("Rs", "").replace("INR", "").replace("₹", "").strip()
#     m = re.search(r"[\d,]+(?:\.\d+)?", s)
#     return m.group(0) if m else s

# def find_price(soup, page_text, jsonld):
#     """Finds the price using a prioritized list of methods."""
#     if jsonld.get("price"):
#         return clean_price(jsonld["price"])

#     for meta_key in ("property", "name"):
#         for meta_value in ("product:price:amount", "og:price:amount", "price"):
#             m = soup.find("meta", attrs={meta_key: meta_value})
#             if m and m.get("content"):
#                 return clean_price(m.get("content"))

#     item_price = soup.find(attrs={"itemprop": "price"})
#     if item_price:
#         return clean_price(item_price.get("content") or item_price.get_text())
    
#     candidate = soup.find(lambda tag: tag.name in ("span", "div") and tag.get("class") and any("price" in c.lower() for c in tag.get("class")))
#     if candidate:
#         return clean_price(candidate.get_text(" ", strip=True))
    
#     match = PRICE_RE.search(page_text)
#     return clean_price(match.group(0)) if match else None

# def find_title(soup, jsonld):
#     if jsonld.get("name"):
#         return clean_whitespace(jsonld["name"])
#     h1 = soup.find("h1")
#     if h1:
#         return clean_whitespace(h1.get_text(" ", strip=True))
#     og_title = soup.find("meta", property="og:title")
#     if og_title and og_title.get("content"):
#         return clean_whitespace(og_title.get("content"))
#     if soup.title and soup.title.string:
#         return clean_whitespace(soup.title.string)
#     return None

# def find_description(soup, page_text, jsonld):
#     if jsonld.get("description"):
#         return clean_whitespace(jsonld["description"])
#     meta = soup.find("meta", {"name": "description"}) or soup.find("meta", property="og:description")
#     if meta and meta.get("content"):
#         return clean_whitespace(meta.get("content"))
#     desc_elem = soup.find(attrs={"itemprop": "description"})
#     if desc_elem:
#         return clean_whitespace(desc_elem.get_text(" ", strip=True))
#     paragraphs = [p.get_text(" ", strip=True) for p in soup.find_all("p")]
#     paragraphs = [clean_whitespace(p) for p in paragraphs if p and len(p) > 50]
#     if paragraphs:
#         paragraphs.sort(key=len, reverse=True)
#         return paragraphs[0]
#     return None

# def find_product_id(soup):
#     """Tries to find a unique product ID or SKU."""
#     for name_attr in ["product-id", "id", "product_id"]:
#         prod_id_input = soup.find("input", attrs={"name": name_attr})
#         if prod_id_input and prod_id_input.get("value"):
#             return prod_id_input.get("value")
#     sku_elem = soup.find(attrs={"itemprop": "sku"})
#     if sku_elem:
#         return sku_elem.get("content") or sku_elem.get_text(strip=True)
#     return None

# def extract_images(soup, jsonld, base_url):
#     """Extracts product image URLs using a prioritized approach."""
#     image_urls = []

#     def normalize_url(url):
#         if not url:
#             return None
#         if url.startswith("//"):
#             url = "https:" + url
#         elif url.startswith(('/', '#')):
#             if base_url:
#                 url = urljoin(base_url, url)
#             else:
#                 return None
#         return url.split("?")[0]

#     if jsonld and "images" in jsonld:
#         image_urls.extend(filter(None, [normalize_url(u) for u in jsonld["images"]]))
#         if image_urls:
#             logging.info("Found %d images via JSON-LD.", len(image_urls))
#             return dedup_preserve_order(image_urls)

#     gallery_selectors = [
#         ".product-gallery", ".product__photos", ".woocommerce-product-gallery",
#         ".product-image-main", ".flickity-slider", ".product-gallery__carousel"
#     ]
#     for selector in gallery_selectors:
#         gallery = soup.select_one(selector)
#         if gallery:
#             for img in gallery.find_all("img"):
#                 src = img.get("src") or img.get("data-src")
#                 image_urls.append(normalize_url(src))
#             if image_urls:
#                 logging.info("Found %d images in gallery container: '%s'", len(image_urls), selector)
#                 return dedup_preserve_order(filter(None, image_urls))

#     for a_tag in soup.find_all("a"):
#         href = a_tag.get("href")
#         if href and any(href.lower().endswith(ext) for ext in [".jpg", ".jpeg", ".png", ".webp"]):
#             if a_tag.find("img"):
#                 image_urls.append(normalize_url(href))
#     if image_urls:
#         logging.info("Found %d images from anchor tags.", len(image_urls))
#         return dedup_preserve_order(filter(None, image_urls))

#     logging.info("Using fallback method: searching for all <img> tags.")
#     for img in soup.find_all("img"):
#         src = img.get("src") or img.get("data-src")
#         image_urls.append(normalize_url(src))
    
#     return dedup_preserve_order(filter(None, image_urls))

# # ---------------------------
# # Main processor
# # ---------------------------
# def process_html_file(file_path, output_folder_path):
#     """Processes a single HTML file and extracts product data."""
#     logging.info("Processing file: %s", os.path.basename(file_path))
#     with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
#         html = f.read()

#     soup = BeautifulSoup(html, "lxml")
#     page_text = soup.get_text("\n", strip=True)

#     jsonld = parse_json_ld(soup)
#     base_url = get_base_url(soup, file_path)

#     data = OrderedDict()
#     data["title"] = find_title(soup, jsonld) or os.path.basename(file_path)
#     data["product_id"] = find_product_id(soup) or os.path.splitext(os.path.basename(file_path))[0]
#     data["price"] = find_price(soup, page_text, jsonld) or "Not found"
#     data["description"] = find_description(soup, page_text, jsonld) or "Not found"
    
#     image_urls = extract_images(soup, jsonld, base_url)
#     data["saved_images"] = save_product_images(data["title"], data["product_id"], image_urls, output_folder_path)
    
#     return data

# def process_html_folder_to_docx(folder_path, output_docx_filename="products.docx"):
#     """Main function to process a folder of HTML files into a DOCX report."""
#     html_files = []
#     for root, _, files in os.walk(folder_path):
#         for fname in files:
#             if fname.lower().endswith((".html", ".htm")):
#                 html_files.append(os.path.join(root, fname))

#     if not html_files:
#         logging.warning("No HTML files found in %s", folder_path)
#         return

#     # --- Create the main timestamped output folder ---
#     timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
#     output_folder_path = os.path.join(os.getcwd(), f"output_{timestamp}")
#     os.makedirs(output_folder_path, exist_ok=True)
#     logging.info(f"Created output directory: {output_folder_path}")

#     doc = Document()
#     doc.add_heading("Extracted Product Details", level=0)

#     for file_path in sorted(html_files):
#         try:
#             # Pass the main output path to the processing function
#             product_data = process_html_file(file_path, output_folder_path)
            
#             doc.add_heading(product_data["title"], level=1)
#             doc.add_paragraph(f"Source file: {os.path.basename(file_path)}")
#             doc.add_paragraph(f"Product ID/SKU: {product_data['product_id']}")
#             doc.add_paragraph(f"Price: {product_data['price']}")
#             doc.add_paragraph("Description:")
#             doc.add_paragraph(product_data["description"])

#             if product_data["saved_images"]:
#                 doc.add_paragraph("Images:")
#                 try:
#                     doc.add_picture(product_data["saved_images"][0], width=Inches(2.5))
#                 except Exception as e:
#                     logging.warning(f"Could not add image {product_data['saved_images'][0]} to docx: {e}")

#             doc.add_page_break()
#         except Exception as e:
#             logging.error(f"Failed to process file {file_path}. Error: {e}")
#             continue

#     # Save the DOCX file inside the timestamped folder
#     final_docx_path = os.path.join(output_folder_path, output_docx_filename)
#     doc.save(final_docx_path)
#     logging.info("Saved combined docx to: %s", final_docx_path)

# # ---------------------------
# # CLI
# # ---------------------------
# def main():
#     parser = argparse.ArgumentParser(description="Extract product details + images from HTML files and write to DOCX.")
#     parser.add_argument("--input-dir", "-i", required=True, help="Folder containing HTML files")
#     parser.add_argument("--output", "-o", default="products.docx", help="Output .docx filename")
#     args = parser.parse_args()

#     if not os.path.isdir(args.input_dir):
#         logging.error("Input directory not found: %s", args.input_dir)
#         return

#     process_html_folder_to_docx(args.input_dir, output_docx_filename=args.output)

# if __name__ == "__main__":
#     main()





























#!/usr/bin/env python3
"""
extract_html_to_docx.py

This script extracts detailed product information (title, price, description,
images, and key features) from a folder of local HTML files. It then generates
a single DOCX report and saves all downloaded product images into a timestamped
output folder.

Usage:
    python extract_html_to_docx.py --input-dir path/to/html_folder --output products.docx

Dependencies:
    pip install beautifulsoup4 lxml python-docx requests Pillow
"""

import os
import re
import json
import argparse
import logging
import requests
import io
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches
from collections import OrderedDict
from PIL import Image
from urllib.parse import urljoin
from datetime import datetime

# ---------------------------
# Logging
# ---------------------------
logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")

# ---------------------------
# Utility / cleaning helpers
# ---------------------------
def clean_whitespace(s: str) -> str:
    """Removes extra whitespace from a string."""
    if not s:
        return s
    return " ".join(s.split())

def dedup_preserve_order(items):
    """Deduplicates items in a list while preserving their original order."""
    seen = set()
    out = []
    for it in items:
        it = it.strip()
        if not it or it in seen:
            continue
        seen.add(it)
        out.append(it)
    return out

def sanitize_filename(name):
    """Removes invalid characters from a string to make it a valid filename."""
    return re.sub(r'[\\/*?:"<>|]', "", name)

def get_base_url(soup, file_path):
    """Extracts the base URL from the page to resolve relative URLs."""
    base_tag = soup.find('base', href=True)
    if base_tag and base_tag['href']:
        return base_tag['href']
    
    og_url = soup.find('meta', property='og:url')
    if og_url and og_url.get('content'):
        return og_url['content']
        
    canonical = soup.find('link', rel='canonical')
    if canonical and canonical.get('href'):
        return canonical['href']
    
    return f"file://{os.path.abspath(file_path)}"

def filter_features(features):
    """Removes common, non-descriptive phrases from the features list."""
    clean = []
    skip_patterns = [
        r"add to cart", r"buy now", r"you save", r"% off",
        r"original price", r"current price", r"regular price",
        r"sale price", r"discount", r"offer",
        r"shop", r"home", r"about", r"contact", r"policy", r"terms", r"faq",
        r"login", r"account", r"blog", r"newsletter",
        r"whatsapp", r"phone", r"email", r"support",
        r"read more", r"media", r"review", r"gallery",
        r"http", r".com", r".in", r".net"
    ]

    for f in features:
        if any(re.search(p, f, re.I) for p in skip_patterns):
            continue
        clean.append(f)
    return clean

# ---------------------------
# Image download logic
# ---------------------------
def download_image(image_url, folder_path, image_name):
    """Download and save an image as JPG with resizing."""
    if not image_url:
        return None
    try:
        response = requests.get(image_url, timeout=20, headers={"User-Agent": "Mozilla/5.0"})
        response.raise_for_status()
        img = Image.open(io.BytesIO(response.content))
        
        if img.size[0] < 640 or img.size[1] < 640:
            resample_filter = getattr(Image, "Resampling", Image).LANCZOS
            img = img.resize((650, 650), resample_filter)
        
        if img.mode != "RGB":
            img = img.convert("RGB")
            
        safe_image_name = sanitize_filename(image_name)
        filename = f"{safe_image_name}.jpg"
        os.makedirs(folder_path, exist_ok=True)
        filepath = os.path.join(folder_path, filename)
        img.save(filepath, "JPEG", quality=95)
        return filepath
    except requests.exceptions.MissingSchema:
        logging.warning(f"Skipping invalid image URL (missing http/https): {image_url}")
        return None
    except Exception as e:
        logging.warning(f"Failed to download image {image_url}. Error: {e}")
        return None

def save_product_images(product_name, product_id, image_urls, base_output_path):
    """Download all images for a product into a dedicated folder inside the base output path."""
    if not product_name and not product_id:
        logging.warning("Skipping image download due to missing product name/ID.")
        return []
    
    product_folder = os.path.join(base_output_path, sanitize_filename(product_id or product_name))
    os.makedirs(product_folder, exist_ok=True)
    saved = []
    
    for idx, url in enumerate(image_urls):
        img_path = download_image(url, product_folder, f"{product_id}_{idx+1}")
        if img_path:
            saved.append(img_path)
    return saved

# ---------------------------
# JSON-LD parsing
# ---------------------------
def parse_json_ld(soup):
    """Parses JSON-LD scripts to find product information."""
    results = {}
    scripts = soup.find_all("script", type="application/ld+json")
    for script in scripts:
        try:
            data = json.loads(script.string or "".join(script.contents))
        except (json.JSONDecodeError, TypeError):
            continue

        if isinstance(data, list):
            for item in data:
                if isinstance(item, dict) and item.get("@type", "").lower() in ("product", "productgroup", "productmodel"):
                    data = item
                    break
        
        if isinstance(data, dict) and data.get("@type", "").lower() in ("product", "productgroup", "productmodel"):
            results["name"] = data.get("name") or results.get("name")
            results["description"] = data.get("description") or results.get("description")
            
            image_data = data.get("image")
            images = []
            if isinstance(image_data, str):
                images.append(image_data)
            elif isinstance(image_data, list):
                for img in image_data:
                    if isinstance(img, str):
                        images.append(img)
                    elif isinstance(img, dict):
                        images.append(img.get("url") or img.get("contentUrl"))
            if images:
                results["images"] = images
            
            offers = data.get("offers", {})
            if isinstance(offers, list):
                offers = offers[0] if offers else {}
            
            if isinstance(offers, dict):
                price = offers.get("price") or offers.get("priceSpecification", {}).get("price")
                currency = offers.get("priceCurrency")
                if price:
                    results["price"] = (str(price) if not currency else f"{price} {currency}")
            
            if "additionalProperty" in data:
                ap = data["additionalProperty"]
                if isinstance(ap, list):
                    props = []
                    for a in ap:
                        name = a.get("name")
                        value = a.get("value")
                        if name and value:
                            props.append(f"{name}: {value}")
                    results.setdefault("additional_properties", []).extend(props)

            if "name" in results:
                return results
                
    return results

# ---------------------------
# Extraction routines
# ---------------------------
PRICE_RE = re.compile(r"(?:₹|Rs\.?|INR)\s?[\d,]+(?:\.\d+)?", flags=re.IGNORECASE)

def clean_price(s):
    if s is None:
        return None
    s = str(s).replace("Rs.", "").replace("Rs", "").replace("INR", "").replace("₹", "").strip()
    m = re.search(r"[\d,]+(?:\.\d+)?", s)
    return m.group(0) if m else s

def find_price(soup, page_text, jsonld):
    if jsonld.get("price"):
        return clean_price(jsonld["price"])
    for meta_key in ("property", "name"):
        for meta_value in ("product:price:amount", "og:price:amount", "price"):
            m = soup.find("meta", attrs={meta_key: meta_value})
            if m and m.get("content"):
                return clean_price(m.get("content"))
    item_price = soup.find(attrs={"itemprop": "price"})
    if item_price:
        return clean_price(item_price.get("content") or item_price.get_text())
    candidate = soup.find(lambda tag: tag.name in ("span", "div") and tag.get("class") and any("price" in c.lower() for c in tag.get("class")))
    if candidate:
        return clean_price(candidate.get_text(" ", strip=True))
    match = PRICE_RE.search(page_text)
    return clean_price(match.group(0)) if match else None

def find_title(soup, jsonld):
    if jsonld.get("name"):
        return clean_whitespace(jsonld["name"])
    h1 = soup.find("h1")
    if h1:
        return clean_whitespace(h1.get_text(" ", strip=True))
    og_title = soup.find("meta", property="og:title")
    if og_title and og_title.get("content"):
        return clean_whitespace(og_title.get("content"))
    if soup.title and soup.title.string:
        return clean_whitespace(soup.title.string)
    return None

def find_description(soup, page_text, jsonld):
    if jsonld.get("description"):
        return clean_whitespace(jsonld["description"])
    meta = soup.find("meta", {"name": "description"}) or soup.find("meta", property="og:description")
    if meta and meta.get("content"):
        return clean_whitespace(meta.get("content"))
    desc_elem = soup.find(attrs={"itemprop": "description"})
    if desc_elem:
        return clean_whitespace(desc_elem.get_text(" ", strip=True))
    paragraphs = [p.get_text(" ", strip=True) for p in soup.find_all("p")]
    paragraphs = [clean_whitespace(p) for p in paragraphs if p and len(p) > 50]
    if paragraphs:
        paragraphs.sort(key=len, reverse=True)
        return paragraphs[0]
    return None

def find_product_id(soup):
    for name_attr in ["product-id", "id", "product_id"]:
        prod_id_input = soup.find("input", attrs={"name": name_attr})
        if prod_id_input and prod_id_input.get("value"):
            return prod_id_input.get("value")
    sku_elem = soup.find(attrs={"itemprop": "sku"})
    if sku_elem:
        return sku_elem.get("content") or sku_elem.get_text(strip=True)
    return None

def extract_images(soup, jsonld, base_url):
    image_urls = []
    def normalize_url(url):
        if not url:
            return None
        if url.startswith("//"):
            url = "https:" + url
        elif url.startswith(('/', '#')):
            if base_url:
                url = urljoin(base_url, url)
            else:
                return None
        return url.split("?")[0]

    if jsonld and "images" in jsonld:
        image_urls.extend(filter(None, [normalize_url(u) for u in jsonld["images"]]))
        if image_urls:
            return dedup_preserve_order(image_urls)
    gallery_selectors = [".product-gallery", ".product__photos", ".woocommerce-product-gallery", ".product-image-main", ".flickity-slider", ".product-gallery__carousel"]
    for selector in gallery_selectors:
        gallery = soup.select_one(selector)
        if gallery:
            for img in gallery.find_all("img"):
                src = img.get("src") or img.get("data-src")
                image_urls.append(normalize_url(src))
            if image_urls:
                return dedup_preserve_order(filter(None, image_urls))
    for a_tag in soup.find_all("a"):
        href = a_tag.get("href")
        if href and any(href.lower().endswith(ext) for ext in [".jpg", ".jpeg", ".png", ".webp"]):
            if a_tag.find("img"):
                image_urls.append(normalize_url(href))
    if image_urls:
        return dedup_preserve_order(filter(None, image_urls))
    for img in soup.find_all("img"):
        src = img.get("src") or img.get("data-src")
        image_urls.append(normalize_url(src))
    return dedup_preserve_order(filter(None, image_urls))

def extract_features(soup, page_text, jsonld):
    """Extracts product features from JSON-LD, lists, tables, and text."""
    features = []
    if jsonld.get("additional_properties"):
        features.extend(jsonld["additional_properties"])
    
    lists = soup.find_all(
        lambda tag: tag.name in ("ul", "ol") and (
            (tag.get("class") and any(re.search(r"(spec|feature|detail|attribute|product)", x, re.I) for x in tag.get("class"))) or
            (tag.get("id") and re.search(r"(spec|feature|detail|attribute|product)", tag.get("id"), re.I))
        )
    )
    for ul in lists:
        for li in ul.find_all("li"):
            t = clean_whitespace(li.get_text(" ", strip=True))
            if t and len(t) > 3:
                features.append(t)
    
    for table in soup.find_all("table"):
        for tr in table.find_all("tr"):
            tds = tr.find_all(["td", "th"])
            if len(tds) >= 2:
                left = clean_whitespace(tds[0].get_text(" ", strip=True))
                right = clean_whitespace(" ".join(td.get_text(" ", strip=True) for td in tds[1:]))
                if left and right:
                    features.append(f"{left}: {right}")
    
    for line in page_text.splitlines():
        line = line.strip()
        if ":" in line:
            left, right = line.split(":", 1)
            if 1 < len(left) < 40 and len(right) > 1 and not re.search(r"(home|about|shop|contact|policy|terms|login|faq|order|account)", left, re.I):
                features.append(f"{left.strip()}: {right.strip()}")

    features = [clean_whitespace(f) for f in features if f and len(f) > 3]
    features = dedup_preserve_order(features)
    features = filter_features(features)
    return features

# ---------------------------
# Main processor
# ---------------------------
def process_html_file(file_path, output_folder_path):
    """Processes a single HTML file and extracts all product data."""
    logging.info("Processing file: %s", os.path.basename(file_path))
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        html = f.read()

    soup = BeautifulSoup(html, "lxml")
    page_text = soup.get_text("\n", strip=True)
    jsonld = parse_json_ld(soup)
    base_url = get_base_url(soup, file_path)

    data = OrderedDict()
    data["title"] = find_title(soup, jsonld) or os.path.basename(file_path)
    data["product_id"] = find_product_id(soup) or os.path.splitext(os.path.basename(file_path))[0]
    data["price"] = find_price(soup, page_text, jsonld) or "Not found"
    data["description"] = find_description(soup, page_text, jsonld) or "Not found"
    data["key_features"] = extract_features(soup, page_text, jsonld) or []
    
    image_urls = extract_images(soup, jsonld, base_url)
    data["saved_images"] = save_product_images(data["title"], data["product_id"], image_urls, output_folder_path)
    
    return data

def process_html_folder_to_docx(folder_path, output_docx_filename="products.docx"):
    """Main function to process a folder of HTML files into a DOCX report."""
    html_files = []
    for root, _, files in os.walk(folder_path):
        for fname in files:
            if fname.lower().endswith((".html", ".htm")):
                html_files.append(os.path.join(root, fname))

    if not html_files:
        logging.warning("No HTML files found in %s", folder_path)
        return

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_folder_path = os.path.join(os.getcwd(), f"output_{timestamp}")
    os.makedirs(output_folder_path, exist_ok=True)
    logging.info(f"Created output directory: {output_folder_path}")

    doc = Document()
    doc.add_heading("Extracted Product Details", level=0)

    for file_path in sorted(html_files):
        try:
            product_data = process_html_file(file_path, output_folder_path)
            
            doc.add_heading(product_data["title"], level=1)
            doc.add_paragraph(f"Source file: {os.path.basename(file_path)}")
            doc.add_paragraph(f"Product ID/SKU: {product_data['product_id']}")
            doc.add_paragraph(f"Price: {product_data['price']}")
            
            doc.add_paragraph("Description:")
            doc.add_paragraph(product_data["description"])

            if product_data["key_features"]:
                doc.add_paragraph("Key Features:")
                for feat in product_data["key_features"]:
                    doc.add_paragraph(feat, style="List Bullet")

            if product_data["saved_images"]:
                doc.add_paragraph("Images:")
                try:
                    doc.add_picture(product_data["saved_images"][0], width=Inches(2.5))
                except Exception as e:
                    logging.warning(f"Could not add image {product_data['saved_images'][0]} to docx: {e}")

            doc.add_page_break()
        except Exception as e:
            logging.error(f"Failed to process file {file_path}. Error: {e}")
            continue

    final_docx_path = os.path.join(output_folder_path, output_docx_filename)
    doc.save(final_docx_path)
    logging.info("Saved combined docx to: %s", final_docx_path)

# ---------------------------
# CLI
# ---------------------------
def main():
    parser = argparse.ArgumentParser(description="Extract product details + images from HTML files and write to DOCX.")
    parser.add_argument("--input-dir", "-i", required=True, help="Folder containing HTML files")
    parser.add_argument("--output", "-o", default="products.docx", help="Output .docx filename")
    args = parser.parse_args()

    if not os.path.isdir(args.input_dir):
        logging.error("Input directory not found: %s", args.input_dir)
        return

    process_html_folder_to_docx(args.input_dir, output_docx_filename=args.output)

if __name__ == "__main__":
    main()


    
